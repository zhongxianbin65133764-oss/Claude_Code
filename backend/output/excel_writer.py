from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from typing import Dict, List, Any
from datetime import datetime
import os


class ExcelWriter:
    """Excel输出器，将提取的数据写入Excel文件"""

    def __init__(self, output_path: str = None):
        if output_path is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_path = f"output_files/extracted_data_{timestamp}.xlsx"
        self.output_path = output_path
        self.workbook = Workbook()
        # 删除默认创建的工作表
        if 'Sheet' in self.workbook.sheetnames:
            del self.workbook['Sheet']

    def _apply_header_style(self, cell):
        """应用表头样式"""
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def _apply_cell_style(self, cell):
        """应用单元格样式"""
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def write_keyword_data(self, keyword_data_list: List[Dict[str, str]]):
        """写入关键词提取的数据"""
        ws = self.workbook.create_sheet("关键信息")

        # 获取所有可能的字段
        all_fields = set()
        for data in keyword_data_list:
            all_fields.update(data.keys())

        headers = ['文档序号'] + sorted(list(all_fields))

        # 写入表头
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            self._apply_header_style(cell)

        # 写入数据
        for row_idx, keyword_data in enumerate(keyword_data_list, 2):
            ws.cell(row=row_idx, column=1, value=f"文档{row_idx-1}")
            self._apply_cell_style(ws.cell(row=row_idx, column=1))

            for col_idx, field in enumerate(headers[1:], 2):
                value = keyword_data.get(field, '')
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                self._apply_cell_style(cell)

        # 调整列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

    def write_table_data(self, table_data: List[Dict[str, Any]]):
        """写入表格数据"""
        if not table_data:
            return

        ws = self.workbook.create_sheet("表格数据")

        # 获取所有字段
        all_fields = set()
        for row in table_data:
            all_fields.update(row.keys())

        headers = sorted(list(all_fields))

        # 写入表头
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            self._apply_header_style(cell)

        # 写入数据
        for row_idx, row_data in enumerate(table_data, 2):
            for col_idx, field in enumerate(headers, 1):
                value = row_data.get(field, '')
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                self._apply_cell_style(cell)

        # 调整列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

    def write_clauses(self, all_clauses: List[str]):
        """写入关键条款"""
        if not all_clauses:
            return

        ws = self.workbook.create_sheet("关键条款")

        # 表头
        headers = ['序号', '条款内容']
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            self._apply_header_style(cell)

        # 写入数据
        for idx, clause in enumerate(all_clauses, 1):
            ws.cell(row=idx+1, column=1, value=idx)
            self._apply_cell_style(ws.cell(row=idx+1, column=1))

            cell = ws.cell(row=idx+1, column=2, value=clause)
            self._apply_cell_style(cell)

        # 调整列宽
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 80

    def write_summary(self, documents_info: List[Dict[str, str]]):
        """写入文档摘要"""
        ws = self.workbook.create_sheet("文档摘要", 0)

        # 表头
        headers = ['序号', '文档类型', '内容预览']
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            self._apply_header_style(cell)

        # 写入数据
        for idx, doc_info in enumerate(documents_info, 1):
            ws.cell(row=idx+1, column=1, value=idx)
            self._apply_cell_style(ws.cell(row=idx+1, column=1))

            ws.cell(row=idx+1, column=2, value=doc_info.get('file_type', ''))
            self._apply_cell_style(ws.cell(row=idx+1, column=2))

            cell = ws.cell(row=idx+1, column=3, value=doc_info.get('preview', ''))
            self._apply_cell_style(cell)

        # 调整列宽
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 60

    def save(self, merged_data: Dict[str, Any]) -> str:
        """保存Excel文件"""
        # 写入各类数据
        self.write_summary(merged_data.get('documents', []))
        self.write_keyword_data(merged_data.get('all_keyword_data', []))
        self.write_table_data(merged_data.get('all_table_data', []))
        self.write_clauses(merged_data.get('all_clauses', []))

        # 确保输出目录存在
        os.makedirs(os.path.dirname(self.output_path), exist_ok=True)

        # 保存文件
        self.workbook.save(self.output_path)
        return self.output_path
