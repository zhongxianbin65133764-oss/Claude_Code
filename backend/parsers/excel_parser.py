from openpyxl import load_workbook
from typing import Dict, List, Any


class ExcelParser:
    """Excel文档解析器"""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def extract_all_sheets(self) -> Dict[str, List[List[Any]]]:
        """提取所有工作表的数据"""
        sheets_data = {}
        try:
            workbook = load_workbook(self.file_path, data_only=True)
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                sheet_data = []
                for row in sheet.iter_rows(values_only=True):
                    sheet_data.append(list(row))
                sheets_data[sheet_name] = sheet_data
        except Exception as e:
            raise Exception(f"Excel解析错误: {str(e)}")
        return sheets_data

    def extract_text(self) -> str:
        """将Excel数据转换为文本格式"""
        text = ""
        sheets_data = self.extract_all_sheets()
        for sheet_name, sheet_data in sheets_data.items():
            text += f"工作表: {sheet_name}\n"
            for row in sheet_data:
                row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                text += row_text + "\n"
            text += "\n"
        return text

    def extract_structured_data(self) -> Dict[str, Any]:
        """提取结构化数据"""
        sheets_data = self.extract_all_sheets()
        text = self.extract_text()

        return {
            'text': text,
            'sheets': sheets_data,
            'file_type': 'excel'
        }
