from docx import Document
from typing import Dict, List, Any


class WordParser:
    """Word文档解析器"""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def extract_text(self) -> str:
        """提取Word文档中的所有文本"""
        text = ""
        try:
            doc = Document(self.file_path)
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
        except Exception as e:
            raise Exception(f"Word文档解析错误: {str(e)}")
        return text

    def extract_tables(self) -> List[List[List[str]]]:
        """提取Word文档中的所有表格"""
        tables = []
        try:
            doc = Document(self.file_path)
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                tables.append(table_data)
        except Exception as e:
            raise Exception(f"Word表格提取错误: {str(e)}")
        return tables

    def extract_structured_data(self) -> Dict[str, Any]:
        """提取结构化数据"""
        text = self.extract_text()
        tables = self.extract_tables()

        return {
            'text': text,
            'tables': tables,
            'file_type': 'word'
        }
